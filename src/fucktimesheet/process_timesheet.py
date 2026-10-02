from playwright.sync_api import sync_playwright
from PIL import Image, ImageDraw, ImageFont
from openpyxl.styles import Alignment, Font
from datetime import datetime, timedelta
from dateutil.relativedelta import relativedelta
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Inches, Pt
import time
import json
import calendar
import copy
import sys
import glob
import os
import pandas as pd
import subprocess
import openpyxl

COORD_2SEN = "3.1356035766711665, 101.69252036618404"
COORD_WFH = "3.195078, 101.758454"


def process_and_convert(oil_days, input_xlsx="timesheet.xlsx", output_dir=r"D:\timesheet", output_pdf="processed_timesheet.pdf"):
    
    matching_files = glob.glob(r"D:\timesheet\*.xls")

    if not matching_files:
        raise FileNotFoundError("No timesheet file found in D:\\timesheet\\")

    latest_file = max(matching_files, key=os.path.getmtime)
    print(f"Processing file: {latest_file}")
    
    input_file = ensure_xlsx(latest_file)
    
    wb = openpyxl.load_workbook(input_file)
    ws = wb.active

    reason_cell = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and "Reason" in str(cell.value):
                reason_cell = cell
                break
        if reason_cell:
            break

    if reason_cell:
        print(f"Found 'Reason' at cell coordinate: {reason_cell.coordinate}")
    else:
        print("Cell containing 'Reason' was not found.")

    create_location_cell(reason_cell, ws)
    populate_reason_col(ws, oil_days)
    populate_location_col(ws)
    
    
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    processed_xlsx = os.path.join(output_dir, "processed_timesheet.xlsx")
    wb.save(processed_xlsx)
    print(f"Successfully generate timesheet at {processed_xlsx}!")

    
    cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", output_dir, processed_xlsx]
    subprocess.run(cmd, check=True)
    expected_pdf = os.path.join(output_dir, "processed_timesheet.pdf")
    print(f"Successfully converted {processed_xlsx} to {expected_pdf}!")

    prepare_signature_page()
    prepare_signature_doc()
    merge_page()

def prepare_signature_doc():
    doc = Document()
    output_dir=r"D:\timesheet"

    section = doc.sections[0]

    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.orientation = WD_ORIENT.LANDSCAPE

    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    doc.add_picture(r"D:\timesheet\assets\ATTENDANCE SIGNATURE BOX (1).png", width=Inches(10.69))

    docx_path = r"D:\timesheet\attendance.signature.docx"
    doc.save(docx_path)

    signature_doc_pdf = os.path.join(output_dir, "attendance.signature.pdf")

    subprocess.run([
        "soffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        output_dir,
        docx_path
    ], check=True)

    print(f"Converted {docx_path} to {signature_doc_pdf}")

def merge_page():
    output_dir = r"D:\timesheet"

    pdfs = [
        os.path.join(output_dir, "processed_timesheet.pdf"),
        os.path.join(output_dir, "attendance.signature.pdf")
    ]

    current_month_year = (datetime.now() - relativedelta(months=1)).strftime("%Y%m")

    merged_pdf = os.path.join(output_dir, f"time_sheet-ahmad_sulha-{current_month_year}.pdf")

    cmd = ["qpdf", "--empty", "--pages"] + pdfs + ["--", merged_pdf]
    subprocess.run(cmd, check=True)
    print(f"Merged pdfs into {merged_pdf}")


def prepare_signature_page():
    signature_page = Image.open(r"D:\timesheet\assets\ATTENDANCE SIGNATURE BOX (1).png")

    erase_own_date(signature_page)
    erase_hm_date(signature_page)
    # erase_signature(signature_page)
    write_own_date(signature_page)
    write_hm_date(signature_page)

    signature_page.save(r"D:\timesheet\assets\ATTENDANCE SIGNATURE BOX (1).png")

def erase_own_date(signature_page):
    draw = ImageDraw.Draw(signature_page)

    x1, y1 = 1400, 150
    x2, y2 = 1920, 220

    draw.rectangle([x1, y1, x2, y2], fill="white")

def erase_hm_date(signature_page):
    draw = ImageDraw.Draw(signature_page)

    x1, y1 = 1393, 473
    x2, y2 = 1916, 554

    draw.rectangle([x1, y1, x2, y2], fill="white")

def erase_signature(signature_page):
    draw = ImageDraw.Draw(signature_page)

    x1, y1 = 468, 110
    x2, y2 = 984, 222

    draw.rectangle([x1, y1, x2, y2], fill="white")

def write_own_date(signature_page):
    draw = ImageDraw.Draw(signature_page)

    font = ImageFont.truetype("arial.ttf", size=36)
    own_date = datetime.now().strftime("01/%m/%Y")

    draw.text((1429, 190), own_date, fill="black", font=font)

def write_hm_date(signature_page):
    draw = ImageDraw.Draw(signature_page)

    font = ImageFont.truetype("arial.ttf", size=36)

    now = datetime.now()
    first_of_month = datetime(now.year, now.month, 1)
    days_until_friday = (4 - first_of_month.weekday()) % 7
    first_friday = first_of_month + timedelta(days=days_until_friday)
    hm_date = first_friday.strftime("%d/%m/%Y")

    draw.text((1429, 520), hm_date, fill="black", font=font)

def ensure_xlsx(file_path: str, output_dir: str = r"D:\timesheet") -> str:
    """Converts .xls to .xlsx natively using LibreOffice, preserving all layout & formatting."""
    if file_path.lower().endswith('.xls'):
        print(f"Converting {file_path} to .xlsx via LibreOffice...")
        
        cmd = ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", output_dir, file_path]
        subprocess.run(cmd, check=True)
        
        # Constructs path: D:\timesheet\filename.xlsx
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        xlsx_path = os.path.join(output_dir, f"{base_name}.xlsx")
        
        return xlsx_path
        
    return file_path

def create_location_cell(cell_to_copy, ws):
    target_cell = ws["BU9"]
    target_cell.value = "Location"

    target_cell.font = copy.copy(cell_to_copy.font)
    target_cell.fill = copy.copy(cell_to_copy.fill)
    target_cell.border = copy.copy(cell_to_copy.border)
    target_cell.alignment = copy.copy(cell_to_copy.alignment)

def populate_location_col(ws):
    font_size_8 = Font(size=8, name="ARIAL")

    for row_num in range(13, 44):
        reason_cell = ws.cell(row=row_num, column=69)
        location_cell = ws[f"BS{row_num}"]

        if reason_cell.value and "wio-2sen" in str(reason_cell.value).lower():
            location_cell.value = COORD_2SEN
            location_cell.font = font_size_8
        elif reason_cell.value and "wfh" in str(reason_cell.value).lower():
            location_cell.value = COORD_WFH
            location_cell.font = font_size_8

def populate_reason_col(ws, oil_days):
    font_size_8 = Font(size=8, name="ARIAL")
    center_align = Alignment(horizontal="center")

    ws.column_dimensions["BR"].width = 25

    oil_days = oil_days.split(",")

    for row_num in range(13, 44):
        day_cell = ws.cell(row=row_num, column=7)
        date_cell = ws.cell(row=row_num, column=1)
        time_in_cell = ws.cell(row=row_num, column=9)
        time_out_cell = ws.cell(row=row_num, column=13)

        day_val = str(day_cell.value or "").lower()
        ordinal_date = str(date_cell.value or "").split('-')[0]
        
        reason_cell = ws[f"BQ{row_num}"]

        ws.merge_cells(f"BQ{row_num}:BR{row_num}")

        if ordinal_date in oil_days:
            reason_cell.value = "OIL"
            reason_cell.font = font_size_8
            reason_cell.alignment = center_align
        elif day_val in ("wed", "thu") and not reason_cell.value:
            reason_cell.value = "WIO-2SEN"
            reason_cell.font = font_size_8
            reason_cell.alignment = center_align
        elif day_val in ("mon", "tue", "fri") and not reason_cell.value:
            reason_cell.value = "WFH"
            reason_cell.font = font_size_8
            reason_cell.alignment = center_align
        elif reason_cell.value and reason_cell.value.lower() == "abs":
            if day_val in ("wed", "thu") and time_in_cell.value is not None and time_in_cell.value == "0":
                reason_cell.value = "WIO-2SEN-Forgot clockin"
                reason_cell.font = font_size_8
                reason_cell.alignment = center_align
            elif day_val in ("wed", "thu") and time_out_cell.value is not None and time_out_cell.value == 0:
                reason_cell.value = "WIO-2SEN-Forgot clockout"
                reason_cell.font = font_size_8
                reason_cell.alignment = center_align
            elif day_val in ("mon", "tue", "fri") and time_in_cell.value is not None and time_in_cell.value == "0":
                reason_cell.value = "WFH-Forgot clockin"
                reason_cell.font = font_size_8
                reason_cell.alignment = center_align
            elif day_val in ("mon", "tue", "fri") and time_out_cell.value is not None and time_out_cell.value == "0":
                reason_cell.value = "WFH-Forgot clockout"
                reason_cell.font = font_size_8
                reason_cell.alignment = center_align

def download_timesheet():
    config_file = os.path.abspath(r"D:\timesheet\assets\config.json")

    print(f"Checking config file at {config_file}")
    if not os.path.exists(config_file):
        raise FileNotFoundError(f"File not found at: {config_file}")

    print(f"Found config file at {config_file}")

    with open(config_file, "r", encoding="utf-8") as f:
        config = json.load(f)

    username = config.get("username")
    password = config.get("password")
    target_dir = os.path.abspath(r"D:\timesheet")
    target_date_from = get_first_day_of_last_month()
    target_date_to = get_last_day_of_last_month()

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, slow_mo=500, args=["--ignore-certificate-errors"])
        context = browser.new_context(accept_downloads=True, ignore_https_errors=True)
        page = context.new_page()

        try:
            page.goto("https://www.infotech-cloudhr.com.my/Login.aspx")
            page.fill('input[name="Credential.UserId"]', username)
            page.fill('input[name="Credential.Password"]', password)
            
            with page.expect_navigation(wait_until="networkidle"):
                page.click('button[type="submit"]')

            print("Clicked login.")

            print("Hovering over top menu 'My Attendance'...")
            page.wait_for_selector("li[data-menuid='12950'] > a", state="visible", timeout=15000)
            page.hover("li[data-menuid='12950'] > a")

            page.get_by_text("My Attendance Report").click()
            print("Clicked My Attendance Report")

            print("Waiting for My Attendance Report page to load")
            page.wait_for_load_state("networkidle")

            print("Entering from date")
            date_selector_from = "input[id='ContentPlaceHolder1_txtFromDate']" 
            page.click(date_selector_from)
            page.fill(date_selector_from, target_date_from)
            page.keyboard.press("Enter")
            print(f"Entered from date as {target_date_from}")

            print("Entering to date")
            date_selector_from = "input[id='ContentPlaceHolder1_txtToDate']" 
            page.click(date_selector_from)
            page.fill(date_selector_from, target_date_to)
            page.keyboard.press("Enter")
            print(f"Entered to date as {target_date_to}")

            print("Clicking Show button to display attendance report sheet")
            page.wait_for_selector("#ContentPlaceHolder1_btnShow", state="visible")

            page.click("#ContentPlaceHolder1_btnShow")
            print("Clicked Show button to display attendance report sheet")

            # wait

            with page.expect_download() as download_info:
                page.click('input[id="ContentPlaceHolder1_btnExcel"]')

            download = download_info.value
            download_path = os.path.join(target_dir, download.suggested_filename)
            download.save_as(download_path)
            print(f"Timesheet downloaded to {download_path}")



            page.wait_for_timeout(10000)

        except Exception as e:
            print(f"Something not right happened when browsing. Error: {e}")
            page.wait_for_timeout(10000)
        
        finally:
            context.close()
            browser.close()

def get_first_day_of_last_month():
    today = datetime.now()

    # Calculate previous month and year
    if today.month == 1:
        prev_month = 12
        year = today.year - 1
    else:
        prev_month = today.month - 1
        year = today.year

    return f"01-{prev_month:02d}-{year}"

def get_last_day_of_last_month():
    today = datetime.now()

    # Calculate previous month and year
    if today.month == 1:
        prev_month = 12
        year = today.year - 1
    else:
        prev_month = today.month - 1
        year = today.year

    # calendar.monthrange returns (first_weekday, num_days_in_month)
    _, last_day = calendar.monthrange(year, prev_month)

    return f"{last_day:02d}-{prev_month:02d}-{year}"

def main():
    start_time = time.time()

    input_xlsx = sys.argv[1] if len(sys.argv) > 1 else None
    oil_days = sys.argv[2] if len(sys.argv) > 2 else None 
    download_timesheet()
    process_and_convert(input_xlsx=input_xlsx, oil_days=oil_days)

    elapsed = round(time.time() - start_time)
    minutes, seconds = divmod(elapsed, 60)

    print(f"Done in {minutes:02d} minutes {seconds:02d} seconds.")

if __name__ == "__main__":
    process_and_convert()